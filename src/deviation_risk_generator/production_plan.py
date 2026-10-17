from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class PlannedBatch:
    date: date
    product: str
    room_class: str


def load_planned_batches(xlsx_path: str | Path) -> list[PlannedBatch]:
    path = Path(xlsx_path)
    suffix = path.suffix.lower()

    if suffix == ".csv":
        with path.open("r", encoding="utf-8", newline="") as csv_file:
            reader = csv.DictReader(csv_file)
            headers = {
                (header or "").strip().lower(): header
                for header in reader.fieldnames or []
            }
            required = {"date", "product", "room_class"}
            if not required.issubset(headers):
                raise ValueError("Production plan must include headers: date, product, room_class")

            batches: list[PlannedBatch] = []
            for row in reader:
                if not row or not any((value or "").strip() for value in row.values()):
                    continue
                raw_date = (row.get(headers["date"]) or "").strip()
                batch_date = date.fromisoformat(raw_date)
                batches.append(
                    PlannedBatch(
                        date=batch_date,
                        product=str(row.get(headers["product"]) or "").strip(),
                        room_class=str(row.get(headers["room_class"]) or "").strip(),
                    )
                )
            return batches

    workbook = load_workbook(filename=str(path), data_only=True)
    sheet = workbook.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in sheet[1]]

    required = {"date", "product", "room_class"}
    if not required.issubset(set(headers)):
        raise ValueError("Production plan must include headers: date, product, room_class")

    idx = {name: headers.index(name) for name in required}
    batches: list[PlannedBatch] = []

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        raw_date = row[idx["date"]]
        if raw_date is None:
            raise ValueError(f"Production plan row {row_number} is missing a date")
        if hasattr(raw_date, "date"):
            batch_date = raw_date.date()
        else:
            batch_date = date.fromisoformat(str(raw_date))

        batches.append(
            PlannedBatch(
                date=batch_date,
                product=str(row[idx["product"]]).strip(),
                room_class=str(row[idx["room_class"]]).strip(),
            )
        )

    return batches
